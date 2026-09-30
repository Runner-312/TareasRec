from dotenv import load_dotenv
load_dotenv()

import os
import io
import re
import uuid
import asyncio
import base64
import hashlib
import hmac
import logging
import time
from collections import deque
from contextlib import asynccontextmanager
from datetime import datetime, timezone, timedelta, date
from zoneinfo import ZoneInfo
from typing import Optional

import jwt
import openpyxl
import requests
from cryptography.fernet import Fernet, InvalidToken
from fastapi import FastAPI, APIRouter, HTTPException, Depends, UploadFile, File, Form, Query, Request
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from fastapi.responses import Response
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_SECRET = os.environ["JWT_SECRET"]
JWT_ALGORITHM = "HS256"
ADMIN_CODE = os.environ.get("ADMIN_CODE")    # solo se usa para crear el admin la primera vez
ADMIN_EMAIL = os.environ.get("ADMIN_EMAIL", "")
TOKEN_DAYS = int(os.environ.get("TOKEN_DAYS", "7"))
FILE_TOKEN_SECONDS = 15 * 60
TRUSTED_PROXY_HOPS = int(os.environ.get("TRUSTED_PROXY_HOPS", "1"))

# Pepper para proteger los PIN en la base de datos. Debe ser distinto de JWT_SECRET
# y NO cambiarse después (si cambia, los PIN guardados dejan de funcionar).
_pepper = os.environ.get("PIN_PEPPER")
if not _pepper or len(_pepper) < 32:
    raise RuntimeError(
        "Falta PIN_PEPPER (mínimo 32 caracteres). Genera uno con: openssl rand -hex 32"
    )
PIN_PEPPER = _pepper.encode()
_fernet = Fernet(base64.urlsafe_b64encode(hashlib.sha256(b"code-enc:" + PIN_PEPPER).digest()))


def hash_code(code: str) -> str:
    return hmac.new(PIN_PEPPER, code.encode(), hashlib.sha256).hexdigest()


def encrypt_code(code: str) -> str:
    return _fernet.encrypt(code.encode()).decode()


def decrypt_code(token: Optional[str]) -> str:
    if not token:
        return ""
    try:
        return _fernet.decrypt(token.encode()).decode()
    except InvalidToken:
        return ""


def public_user(u: dict) -> dict:
    """Quita los campos internos del PIN; el admin sigue viendo el código en 'code'."""
    out = {k: v for k, v in u.items() if k not in ("code_hash", "code_enc", "code")}
    out["code"] = decrypt_code(u.get("code_enc"))
    return out


class AttemptLimiter:
    """Limita intentos fallidos en memoria (una sola instancia del backend)."""

    def __init__(self, max_fail: int, window: int):
        self.max_fail, self.window = max_fail, window
        self.fails = {}

    def _prune(self, key):
        q = self.fails.setdefault(key, deque())
        limit = time.time() - self.window
        while q and q[0] < limit:
            q.popleft()
        return q

    def check(self, key):
        q = self._prune(key)
        if len(q) >= self.max_fail:
            wait = int(q[0] + self.window - time.time()) + 1
            raise HTTPException(
                status_code=429,
                detail=f"Demasiados intentos. Intenta de nuevo en {max(1, wait // 60 + 1)} min",
                headers={"Retry-After": str(max(1, wait))},
            )

    def fail(self, key):
        self._prune(key).append(time.time())

    def reset(self, key):
        self.fails.pop(key, None)


ip_limiter = AttemptLimiter(max_fail=5, window=15 * 60)        # por IP
global_limiter = AttemptLimiter(max_fail=60, window=15 * 60)   # respaldo contra rotación de IPs
change_code_limiter = AttemptLimiter(max_fail=5, window=15 * 60)


def client_ip(request: Request) -> str:
    xff = [p.strip() for p in request.headers.get("x-forwarded-for", "").split(",") if p.strip()]
    if len(xff) >= TRUSTED_PROXY_HOPS > 0:
        return xff[-TRUSTED_PROXY_HOPS]
    return request.client.host if request.client else "unknown"


def sniff_image(data: bytes):
    """Devuelve (content_type, extensión) según los bytes reales, o None si no es imagen permitida."""
    if data[:3] == b"\xff\xd8\xff":
        return "image/jpeg", "jpg"
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "image/png", "png"
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return "image/gif", "gif"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp", "webp"
    if data[4:8] == b"ftyp" and data[8:12] in (b"heic", b"heix", b"hevc", b"heim", b"heis", b"mif1", b"msf1"):
        return "image/heic", "heic"
    return None
APP_NAME = os.environ.get("APP_NAME", "tareasrec")

STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
EMERGENT_KEY = os.environ.get("EMERGENT_LLM_KEY")
storage_key = None


def init_storage(force: bool = False):
    global storage_key
    if storage_key and not force:
        return storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_KEY}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key


def put_object(path: str, data: bytes, content_type: str) -> dict:
    key = init_storage()
    resp = requests.put(
        f"{STORAGE_URL}/objects/{path}",
        headers={"X-Storage-Key": key, "Content-Type": content_type},
        data=data, timeout=120,
    )
    resp.raise_for_status()
    return resp.json()


def get_object(path: str):
    key = init_storage()
    resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    resp.raise_for_status()
    return resp.content, resp.headers.get("Content-Type", "application/octet-stream")


# (app se crea más abajo con lifespan)
api_router = APIRouter(prefix="/api")
bearer = HTTPBearer(auto_error=False)

DAY_LABELS = ["Mié", "Jue", "Vie", "Sáb", "Dom", "Lun", "Mar"]
MONTH_SHORT = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]


def fmt_short(d: date) -> str:
    return f"{d.day} {MONTH_SHORT[d.month - 1]}"
BONUS = {1: 0.30, 2: 0.20, 3: 0.10}
RATE_PER_HOUR = 0.30
GOAL_MINUTES = 600
PAYOUT_CURRENCY = "USDT"
TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "binance_template.xlsx")
TZ = ZoneInfo("America/Caracas")
DEFAULT_RATES = {"kgen_rate": 4.0, "bonus_rate": RATE_PER_HOUR}


def now_local() -> datetime:
    return datetime.now(TZ)


def today_local() -> date:
    return now_local().date()


def week_start_of(d: date) -> date:
    return d - timedelta(days=(d.weekday() - 2) % 7)


async def rates_for_week(ws: date) -> dict:
    doc = await db.rates.find({"week_start": {"$lte": ws.isoformat()}}, {"_id": 0}).sort("week_start", -1).limit(1).to_list(1)
    if doc:
        return {"kgen_rate": float(doc[0]["kgen_rate"]), "bonus_rate": float(doc[0]["bonus_rate"]), "week_start": doc[0]["week_start"]}
    return {**DEFAULT_RATES, "week_start": None}


def calc_payment(minutes: int, rank: Optional[int], bonus_rate: float = RATE_PER_HOUR):
    hours = minutes / 60
    qualifies = hours > 10
    base = round(hours * bonus_rate, 2) if qualifies else 0.0
    bonus_pct = BONUS.get(rank, 0.0) if qualifies and rank else 0.0
    bonus = round(base * bonus_pct, 2)
    return {
        "hours": round(hours, 2),
        "qualifies": qualifies,
        "base": base,
        "bonus_pct": bonus_pct,
        "bonus": bonus,
        "total": round(base + bonus, 2),
        "bonus_rate": bonus_rate,
    }


def create_token(user_id: str, role: str) -> str:
    payload = {
        "sub": user_id,
        "role": role,
        "scope": "auth",
        "exp": datetime.now(timezone.utc) + timedelta(days=TOKEN_DAYS),
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)


def create_file_token(user_id: str) -> str:
    payload = {
        "sub": user_id,
        "scope": "file",
        "exp": datetime.now(timezone.utc) + timedelta(seconds=FILE_TOKEN_SECONDS),
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)


async def get_current_user(creds: HTTPAuthorizationCredentials = Depends(bearer)):
    if not creds:
        raise HTTPException(status_code=401, detail="No autenticado")
    try:
        payload = jwt.decode(creds.credentials, JWT_SECRET, algorithms=[JWT_ALGORITHM])
    except jwt.PyJWTError:
        raise HTTPException(status_code=401, detail="Token inválido")
    if payload.get("scope", "auth") != "auth":
        raise HTTPException(status_code=401, detail="Token inválido")
    user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0})
    if not user or not user.get("active", True):
        raise HTTPException(status_code=401, detail="Usuario no encontrado")
    return user


async def require_admin(user=Depends(get_current_user)):
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Solo administrador")
    return user


async def build_week_table(ws: date):
    we = ws + timedelta(days=6)
    workers = await db.users.find({"role": "worker"}, {"_id": 0}).to_list(1000)
    entries = await db.entries.find(
        {"date": {"$gte": ws.isoformat(), "$lte": we.isoformat()}}, {"_id": 0}
    ).to_list(100000)
    minutes_by = {}
    for e in entries:
        minutes_by[e["worker_id"]] = minutes_by.get(e["worker_id"], 0) + e["minutes"]
    rows = []
    for w in workers:
        rows.append({"id": w["id"], "name": w["name"], "minutes": minutes_by.get(w["id"], 0)})
    rows.sort(key=lambda r: (-r["minutes"], r["name"].lower()))
    rates = await rates_for_week(ws)
    for i, r in enumerate(rows):
        r["rank"] = i + 1 if r["minutes"] > 0 else None
        r.update(calc_payment(r["minutes"], r["rank"], rates["bonus_rate"]))
        r["kgen"] = round(r["minutes"] / 60 * rates["kgen_rate"], 2)
    paid = await db.payments.find({"week_start": ws.isoformat()}, {"_id": 0}).to_list(10000)
    paid_map = {p["worker_id"]: p for p in paid}
    for r in rows:
        r["paid"] = r["id"] in paid_map
    return rows, entries


async def build_global_ranking():
    workers = await db.users.find({"role": "worker"}, {"_id": 0}).to_list(1000)
    pipeline = [{"$group": {"_id": "$worker_id", "minutes": {"$sum": "$minutes"}}}]
    agg = await db.entries.aggregate(pipeline).to_list(10000)
    minutes_by = {a["_id"]: a["minutes"] for a in agg}
    rows = [
        {"id": w["id"], "name": w["name"], "minutes": minutes_by.get(w["id"], 0) + int(w.get("historical_minutes") or 0)}
        for w in workers
    ]
    rows.sort(key=lambda r: (-r["minutes"], r["name"].lower()))
    for i, r in enumerate(rows):
        r["rank"] = i + 1 if r["minutes"] > 0 else None
        r["hours"] = round(r["minutes"] / 60, 2)
    return rows


class LoginBody(BaseModel):
    code: str


class WorkerBody(BaseModel):
    name: str
    code: Optional[str] = None
    binance_pay_id: Optional[str] = None
    usdt_bep20_address: Optional[str] = None
    historical_minutes: Optional[int] = None


class WalletBody(BaseModel):
    usdt_bep20_address: str


def clean_wallet(value: Optional[str]) -> Optional[str]:
    if value is None:
        return None
    v = value.strip()
    if v == "":
        return ""
    if not re.fullmatch(r"[0-9A-Za-z]{20,64}", v):
        raise HTTPException(status_code=400, detail="La dirección USDT (BEP20) no parece válida (20-64 letras y números)")
    return v


def clean_binance(value: Optional[str]) -> Optional[str]:
    if value is None:
        return None
    return value.strip()[:40]


class PaymentBody(BaseModel):
    worker_id: str
    week_start: str


class RatesBody(BaseModel):
    kgen_rate: float
    bonus_rate: float


class ChangeCodeBody(BaseModel):
    current_code: str
    new_code: str


class ExtraBonusBody(BaseModel):
    worker_id: str
    week_start: str
    received: Optional[bool] = None
    paid: Optional[bool] = None
    amount: Optional[float] = None


@api_router.post("/admin/change-code")
async def change_admin_code(body: ChangeCodeBody, admin=Depends(require_admin)):
    current = body.current_code.strip()
    new = body.new_code.strip()
    change_code_limiter.check(admin["id"])
    if not hmac.compare_digest(hash_code(current), admin.get("code_hash") or ""):
        change_code_limiter.fail(admin["id"])
        raise HTTPException(status_code=401, detail="El código actual no es correcto")
    change_code_limiter.reset(admin["id"])
    if not (new.isdigit() and len(new) == 4):
        raise HTTPException(status_code=400, detail="El nuevo código debe tener 4 dígitos")
    if new == current:
        raise HTTPException(status_code=400, detail="El nuevo código es igual al actual")
    if await db.users.find_one({"code_hash": hash_code(new), "id": {"$ne": admin["id"]}}):
        raise HTTPException(status_code=409, detail="Ese código ya lo usa un miembro")
    await db.users.update_one({"id": admin["id"]}, {"$set": {"code_hash": hash_code(new), "code_enc": encrypt_code(new), "code_changed_at": datetime.now(timezone.utc).isoformat()}})
    return {"ok": True}


@api_router.get("/admin/extra-bonus")
async def extra_bonus_list(week_start: Optional[str] = None, admin=Depends(require_admin)):
    ws = week_start_of(today_local()) - timedelta(days=7)
    if week_start:
        try:
            ws = week_start_of(date.fromisoformat(week_start))
        except ValueError:
            raise HTTPException(status_code=400, detail="Fecha inválida")
    rows, _ = await build_week_table(ws)
    workers = {w["id"]: w for w in await db.users.find({"role": "worker"}, {"_id": 0}).to_list(1000)}
    marks = {m["worker_id"]: m for m in await db.extra_bonus.find({"week_start": ws.isoformat()}, {"_id": 0}).to_list(1000)}
    items = []
    for r in rows:
        if r["minutes"] <= 300:
            continue
        w = workers.get(r["id"], {})
        m = marks.get(r["id"], {})
        items.append({
            "worker_id": r["id"],
            "name": r["name"],
            "minutes": r["minutes"],
            "hours": r["hours"],
            "member_done": ws.isoformat() in (w.get("weekly_report_done") or []),
            "received": bool(m.get("received")),
            "paid": bool(m.get("paid")),
            "amount": float(m.get("amount") or 0),
            "binance_pay_id": w.get("binance_pay_id") or "",
        })
    return {"week_start": ws.isoformat(), "week_end": (ws + timedelta(days=6)).isoformat(), "items": items}


@api_router.post("/admin/extra-bonus")
async def extra_bonus_mark(body: ExtraBonusBody, admin=Depends(require_admin)):
    try:
        ws = week_start_of(date.fromisoformat(body.week_start)).isoformat()
    except ValueError:
        raise HTTPException(status_code=400, detail="Fecha inválida")
    updates = {"updated_at": datetime.now(timezone.utc).isoformat()}
    if body.received is not None:
        updates["received"] = body.received
    if body.paid is not None:
        updates["paid"] = body.paid
    if body.amount is not None:
        if body.amount < 0:
            raise HTTPException(status_code=400, detail="Monto inválido")
        updates["amount"] = round(body.amount, 2)
    await db.extra_bonus.update_one({"worker_id": body.worker_id, "week_start": ws}, {"$set": updates}, upsert=True)
    doc = await db.extra_bonus.find_one({"worker_id": body.worker_id, "week_start": ws}, {"_id": 0})
    return doc


@api_router.get("/admin/rates")
async def get_rates(admin=Depends(require_admin)):
    ws = week_start_of(today_local())
    current = await rates_for_week(ws)
    history = await db.rates.find({}, {"_id": 0}).sort("week_start", -1).to_list(52)
    return {"current": current, "week_start": ws.isoformat(), "history": history}


@api_router.put("/admin/rates")
async def set_rates(body: RatesBody, admin=Depends(require_admin)):
    if body.kgen_rate < 0 or body.bonus_rate < 0:
        raise HTTPException(status_code=400, detail="Las tarifas no pueden ser negativas")
    ws = week_start_of(today_local()).isoformat()
    doc = {"week_start": ws, "kgen_rate": round(body.kgen_rate, 4), "bonus_rate": round(body.bonus_rate, 4), "updated_at": datetime.now(timezone.utc).isoformat()}
    await db.rates.update_one({"week_start": ws}, {"$set": doc}, upsert=True)
    return doc


@api_router.get("/")
async def root():
    return {"message": "TareasREC API"}


@api_router.post("/auth/login")
async def login(body: LoginBody, request: Request):
    ip = client_ip(request)
    ip_limiter.check(ip)
    global_limiter.check("global")
    code = body.code.strip()
    if not (code.isdigit() and len(code) == 4):
        raise HTTPException(status_code=400, detail="El código debe tener 4 dígitos")
    user = await db.users.find_one({"code_hash": hash_code(code), "active": True}, {"_id": 0})
    if not user:
        ip_limiter.fail(ip)
        global_limiter.fail("global")
        logger.warning("Login fallido desde %s", ip)
        raise HTTPException(status_code=401, detail="Código incorrecto")
    ip_limiter.reset(ip)
    token = create_token(user["id"], user["role"])
    return {"token": token, "user": {"id": user["id"], "name": user["name"], "role": user["role"]}}


@api_router.post("/auth/file-token")
async def file_token(user=Depends(get_current_user)):
    return {"token": create_file_token(user["id"]), "expires_in": FILE_TOKEN_SECONDS}


@api_router.get("/auth/me")
async def me(user=Depends(get_current_user)):
    return {"id": user["id"], "name": user["name"], "role": user["role"]}


@api_router.get("/admin/workers")
async def list_workers(admin=Depends(require_admin)):
    workers = await db.users.find({"role": "worker"}, {"_id": 0}).sort("name", 1).to_list(1000)
    today = today_local()
    ws = week_start_of(today)
    we = ws + timedelta(days=6)
    entries = await db.entries.find(
        {"date": {"$gte": ws.isoformat(), "$lte": we.isoformat()}}, {"_id": 0}
    ).to_list(100000)
    week_min = {}
    for e in entries:
        week_min[e["worker_id"]] = week_min.get(e["worker_id"], 0) + e["minutes"]
    pipeline = [{"$group": {"_id": "$worker_id", "minutes": {"$sum": "$minutes"}}}]
    agg = await db.entries.aggregate(pipeline).to_list(10000)
    total_min = {a["_id"]: a["minutes"] for a in agg}
    workers = [public_user(w) for w in workers]
    for w in workers:
        w["week_minutes"] = week_min.get(w["id"], 0)
        w["historical_minutes"] = int(w.get("historical_minutes") or 0)
        w["binance_pay_id"] = w.get("binance_pay_id") or ""
        w["usdt_bep20_address"] = w.get("usdt_bep20_address") or ""
        w["total_minutes"] = total_min.get(w["id"], 0) + w["historical_minutes"]
    return workers


@api_router.post("/admin/workers")
async def create_worker(body: WorkerBody, admin=Depends(require_admin)):
    name = body.name.strip()
    code = (body.code or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="El nombre es obligatorio")
    if not (code.isdigit() and len(code) == 4):
        raise HTTPException(status_code=400, detail="El código debe tener 4 dígitos")
    existing = await db.users.find_one({"code_hash": hash_code(code)})
    if existing:
        raise HTTPException(status_code=400, detail="Ese código ya está en uso")
    doc = {
        "id": str(uuid.uuid4()),
        "name": name,
        "code_hash": hash_code(code),
        "code_enc": encrypt_code(code),
        "role": "worker",
        "active": True,
        "binance_pay_id": clean_binance(body.binance_pay_id) or "",
        "usdt_bep20_address": clean_wallet(body.usdt_bep20_address) or "",
        "historical_minutes": max(0, int(body.historical_minutes or 0)),
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.users.insert_one(doc)
    doc.pop("_id", None)
    return public_user(doc)


@api_router.put("/admin/workers/{worker_id}")
async def update_worker(worker_id: str, body: WorkerBody, admin=Depends(require_admin)):
    worker = await db.users.find_one({"id": worker_id, "role": "worker"})
    if not worker:
        raise HTTPException(status_code=404, detail="Miembro no encontrado")
    updates = {}
    name = body.name.strip()
    if name:
        updates["name"] = name
    if body.code:
        code = body.code.strip()
        if not (code.isdigit() and len(code) == 4):
            raise HTTPException(status_code=400, detail="El código debe tener 4 dígitos")
        existing = await db.users.find_one({"code_hash": hash_code(code), "id": {"$ne": worker_id}})
        if existing:
            raise HTTPException(status_code=400, detail="Ese código ya está en uso")
        updates["code_hash"] = hash_code(code)
        updates["code_enc"] = encrypt_code(code)
    if body.binance_pay_id is not None:
        updates["binance_pay_id"] = clean_binance(body.binance_pay_id)
    if body.usdt_bep20_address is not None:
        updates["usdt_bep20_address"] = clean_wallet(body.usdt_bep20_address)
    if body.historical_minutes is not None:
        if body.historical_minutes < 0:
            raise HTTPException(status_code=400, detail="Los minutos históricos no pueden ser negativos")
        updates["historical_minutes"] = int(body.historical_minutes)
    if updates:
        await db.users.update_one({"id": worker_id}, {"$set": updates})
    return public_user(await db.users.find_one({"id": worker_id}, {"_id": 0}))


@api_router.delete("/admin/workers/{worker_id}")
async def delete_worker(worker_id: str, admin=Depends(require_admin)):
    result = await db.users.delete_one({"id": worker_id, "role": "worker"})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Miembro no encontrado")
    return {"deleted": True}


@api_router.get("/admin/week")
async def admin_week(start: Optional[str] = None, admin=Depends(require_admin)):
    if start:
        try:
            ws = date.fromisoformat(start)
        except ValueError:
            raise HTTPException(status_code=400, detail="Fecha inválida")
        ws = week_start_of(ws)
    else:
        ws = week_start_of(today_local())
    we = ws + timedelta(days=6)
    payday = ws + timedelta(days=13)
    rows, entries = await build_week_table(ws)
    days = []
    for i in range(7):
        d = ws + timedelta(days=i)
        tot = sum(e["minutes"] for e in entries if e["date"] == d.isoformat())
        days.append({"date": d.isoformat(), "label": DAY_LABELS[i], "total": tot})
    workers_chart = []
    for r in rows:
        data = []
        for i in range(7):
            d = (ws + timedelta(days=i)).isoformat()
            data.append(sum(e["minutes"] for e in entries if e["worker_id"] == r["id"] and e["date"] == d))
        workers_chart.append({"id": r["id"], "name": r["name"], "data": data})
    total_minutes = sum(d["total"] for d in days)
    return {
        "week_start": ws.isoformat(),
        "week_end": we.isoformat(),
        "payday": payday.isoformat(),
        "is_current": ws == week_start_of(today_local()),
        "days": days,
        "workers_chart": workers_chart,
        "table": rows,
        "totals": {
            "minutes": total_minutes,
            "hours": round(total_minutes / 60, 2),
            "avg_per_day": round(total_minutes / 7, 1),
            "active_workers": sum(1 for r in rows if r["minutes"] > 0),
        },
    }


@api_router.get("/admin/rankings")
async def admin_rankings(admin=Depends(require_admin)):
    ws = week_start_of(today_local())
    rows, _ = await build_week_table(ws)
    weekly = [
        {"rank": r["rank"], "id": r["id"], "name": r["name"], "minutes": r["minutes"], "hours": r["hours"]}
        for r in rows
    ]
    global_ranking = await build_global_ranking()
    return {"weekly": weekly, "global": global_ranking}


@api_router.get("/admin/overview")
async def admin_overview(admin=Depends(require_admin)):
    return await build_overview()


async def build_overview():
    today = today_local()
    entries = await db.entries.find({}, {"_id": 0}).to_list(100000)
    workers = {w["id"]: w for w in await db.users.find({"role": "worker"}, {"_id": 0}).to_list(1000)}
    groups = {}
    for e in entries:
        ws = week_start_of(date.fromisoformat(e["date"]))
        if ws + timedelta(days=6) < today:
            key = (e["worker_id"], ws)
            groups[key] = groups.get(key, 0) + e["minutes"]
    weeks = {}
    for (wid, ws), m in groups.items():
        weeks.setdefault(ws, []).append((wid, m))
    rank_of = {}
    for ws, lst in weeks.items():
        lst.sort(key=lambda x: -x[1])
        for i, (wid, m) in enumerate(lst):
            rank_of[(wid, ws)] = i + 1
    paid = await db.payments.find({}, {"_id": 0}).to_list(10000)
    paid_keys = {(p["worker_id"], p["week_start"]) for p in paid}
    total_paid = round(sum(p["amount"] for p in paid), 2)
    rate_cache = {}
    pending = []
    for (wid, ws), m in groups.items():
        if (wid, ws.isoformat()) in paid_keys:
            continue
        if ws not in rate_cache:
            rate_cache[ws] = await rates_for_week(ws)
        calc = calc_payment(m, rank_of[(wid, ws)], rate_cache[ws]["bonus_rate"])
        if calc["total"] > 0:
            pending.append({
                "worker_id": wid,
                "name": workers.get(wid, {}).get("name", "Eliminado"),
                "binance_pay_id": workers.get(wid, {}).get("binance_pay_id", ""),
                "week_start": ws.isoformat(),
                "week_end": (ws + timedelta(days=6)).isoformat(),
                "minutes": m,
                "hours": calc["hours"],
                "rank": rank_of[(wid, ws)],
                "base": calc["base"],
                "bonus": calc["bonus"],
                "total": calc["total"],
            })
    pending.sort(key=lambda p: p["week_start"], reverse=True)
    ws_cur = week_start_of(today)
    we_cur = ws_cur + timedelta(days=6)
    cur_minutes = sum(e["minutes"] for e in entries if ws_cur.isoformat() <= e["date"] <= we_cur.isoformat())
    return {
        "total_paid": total_paid,
        "pending_total": round(sum(p["total"] for p in pending), 2),
        "pending": pending,
        "current_week": {"minutes": cur_minutes, "hours": round(cur_minutes / 60, 2)},
    }


@api_router.post("/admin/payments")
async def mark_paid(body: PaymentBody, admin=Depends(require_admin)):
    return await mark_week_paid(body.worker_id, body.week_start)


async def mark_week_paid(worker_id: str, week_start: str):
    try:
        ws = week_start_of(date.fromisoformat(week_start))
    except ValueError:
        raise HTTPException(status_code=400, detail="Fecha inválida")
    rows, _ = await build_week_table(ws)
    row = next((r for r in rows if r["id"] == worker_id), None)
    if row is None:
        entries = await db.entries.find({"worker_id": worker_id}, {"_id": 0}).to_list(100000)
        we = ws + timedelta(days=6)
        m = sum(e["minutes"] for e in entries if ws.isoformat() <= e["date"] <= we.isoformat())
        calc = calc_payment(m, None, (await rates_for_week(ws))["bonus_rate"])
        row = {"rank": None, **calc}
    doc = {
        "id": str(uuid.uuid4()),
        "worker_id": worker_id,
        "week_start": ws.isoformat(),
        "amount": row["total"],
        "base": row["base"],
        "bonus": row["bonus"],
        "rank": row["rank"],
        "paid_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.payments.update_one(
        {"worker_id": worker_id, "week_start": ws.isoformat()},
        {"$set": doc},
        upsert=True,
    )
    return doc


def payout_rows(pending):
    by_worker = {}
    for p in pending:
        r = by_worker.setdefault(p["worker_id"], {"worker_id": p["worker_id"], "name": p["name"], "binance_pay_id": p["binance_pay_id"], "total": 0.0, "weeks": []})
        r["total"] = round(r["total"] + p["total"], 2)
        r["weeks"].append(p["week_start"])
    rows = sorted(by_worker.values(), key=lambda r: r["name"].lower())
    for r in rows:
        r["ready"] = bool(r["binance_pay_id"]) and r["total"] >= 0.5
        r["issue"] = None if r["ready"] else ("Sin Binance Pay ID" if not r["binance_pay_id"] else "Monto menor a 0.50 USDT")
    return rows


@api_router.get("/admin/payments/export")
async def payout_preview(admin=Depends(require_admin)):
    overview = await build_overview()
    rows = payout_rows(overview["pending"])
    ready = [r for r in rows if r["ready"]]
    return {
        "rows": rows,
        "ready_count": len(ready),
        "ready_total": round(sum(r["total"] for r in ready), 2),
        "currency": PAYOUT_CURRENCY,
    }


def fill_binance_template(rows):
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb.active
    for i, r in enumerate(rows[:250]):
        row = 3 + i
        ws.cell(row=row, column=1, value=r["account_type"])
        ws.cell(row=row, column=2, value=r["account"])
        ws.cell(row=row, column=3, value=r["currency"])
        ws.cell(row=row, column=4, value=r["amount"])
        ws.cell(row=row, column=5, value=r.get("note") or "")
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="binance_pay_{today_local().isoformat()}.xlsx"'},
    )


@api_router.get("/admin/payments/export.xlsx")
async def payout_xlsx(admin=Depends(require_admin)):
    overview = await build_overview()
    rows = [r for r in payout_rows(overview["pending"]) if r["ready"]]
    if not rows:
        raise HTTPException(status_code=400, detail="No hay pagos listos para exportar")
    return fill_binance_template([
        {"account_type": "Binance ID (BUID)", "account": r["binance_pay_id"], "currency": PAYOUT_CURRENCY, "amount": r["total"], "note": r["name"]}
        for r in rows
    ])


ACCOUNT_TYPES = ("Binance ID (BUID)", "Binance Registered Email")


class PayoutRowBody(BaseModel):
    account_type: str
    account: str
    currency: str = PAYOUT_CURRENCY
    amount: float
    note: Optional[str] = ""


class CustomPayoutBody(BaseModel):
    rows: list[PayoutRowBody]


@api_router.post("/admin/payments/custom.xlsx")
async def custom_payout_xlsx(body: CustomPayoutBody, admin=Depends(require_admin)):
    if not body.rows:
        raise HTTPException(status_code=400, detail="Agrega al menos un destinatario")
    if len(body.rows) > 250:
        raise HTTPException(status_code=400, detail="Máximo 250 destinatarios por archivo")
    clean = []
    for i, r in enumerate(body.rows, start=1):
        if r.account_type not in ACCOUNT_TYPES:
            raise HTTPException(status_code=400, detail=f"Fila {i}: tipo de cuenta inválido")
        account = r.account.strip()
        if not account:
            raise HTTPException(status_code=400, detail=f"Fila {i}: falta el ID o correo")
        if r.amount <= 0:
            raise HTTPException(status_code=400, detail=f"Fila {i}: el monto debe ser mayor a 0")
        clean.append({"account_type": r.account_type, "account": account, "currency": (r.currency or PAYOUT_CURRENCY).strip().upper()[:10], "amount": round(r.amount, 2), "note": (r.note or "")[:60]})
    return fill_binance_template(clean)


@api_router.post("/admin/payments/mark-all")
async def mark_all_paid(admin=Depends(require_admin)):
    overview = await build_overview()
    ready_ids = {r["worker_id"] for r in payout_rows(overview["pending"]) if r["ready"]}
    count = 0
    for p in overview["pending"]:
        if p["worker_id"] in ready_ids:
            await mark_week_paid(p["worker_id"], p["week_start"])
            count += 1
    return {"marked": count}


@api_router.delete("/admin/payments")
async def unmark_paid(worker_id: str = Query(...), week_start: str = Query(...), admin=Depends(require_admin)):
    ws = week_start_of(date.fromisoformat(week_start))
    await db.payments.delete_one({"worker_id": worker_id, "week_start": ws.isoformat()})
    return {"deleted": True}


@api_router.get("/admin/entries")
async def admin_entries(
    worker_id: Optional[str] = None,
    week_start: Optional[str] = None,
    admin=Depends(require_admin),
):
    query = {}
    if worker_id:
        query["worker_id"] = worker_id
    if week_start:
        try:
            ws = week_start_of(date.fromisoformat(week_start))
        except ValueError:
            raise HTTPException(status_code=400, detail="Fecha inválida")
        query["date"] = {"$gte": ws.isoformat(), "$lte": (ws + timedelta(days=6)).isoformat()}
    entries = await db.entries.find(query, {"_id": 0}).sort("date", -1).to_list(500)
    workers = {w["id"]: w["name"] for w in await db.users.find({"role": "worker"}, {"_id": 0}).to_list(1000)}
    for e in entries:
        e["worker_name"] = workers.get(e["worker_id"], "Eliminado")
        e["reviewed"] = bool(e.get("reviewed"))
    return entries


@api_router.patch("/admin/entries/{entry_id}/review")
async def review_entry(entry_id: str, reviewed: bool = Query(True), admin=Depends(require_admin)):
    result = await db.entries.update_one({"id": entry_id}, {"$set": {"reviewed": reviewed}})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Registro no encontrado")
    return {"id": entry_id, "reviewed": reviewed}


@api_router.delete("/admin/entries/{entry_id}")
async def delete_entry(entry_id: str, admin=Depends(require_admin)):
    result = await db.entries.delete_one({"id": entry_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Registro no encontrado")
    return {"deleted": True}


@api_router.post("/me/wallet")
async def set_my_wallet(body: WalletBody, user=Depends(get_current_user)):
    if user.get("role") != "worker":
        raise HTTPException(status_code=403, detail="Solo miembros")
    if user.get("usdt_bep20_address"):
        raise HTTPException(status_code=400, detail="Ya registraste tu dirección; pídele a Wuilber si necesitas cambiarla")
    addr = clean_wallet(body.usdt_bep20_address)
    if not addr:
        raise HTTPException(status_code=400, detail="Ingresa tu dirección USDT (BEP20)")
    await db.users.update_one({"id": user["id"]}, {"$set": {"usdt_bep20_address": addr}})
    return {"usdt_bep20_address": addr}


@api_router.post("/entries")
async def create_entry(
    minutes: int = Form(...),
    screenshot: UploadFile = File(...),
    date_str: Optional[str] = Form(None, alias="date"),
    user=Depends(get_current_user),
):
    if user.get("role") != "worker":
        raise HTTPException(status_code=403, detail="Solo miembros")
    if minutes <= 0 or minutes > 1440:
        raise HTTPException(status_code=400, detail="Minutos inválidos")
    today = today_local()
    target = today
    if date_str:
        try:
            target = date.fromisoformat(date_str)
        except ValueError:
            raise HTTPException(status_code=400, detail="Fecha inválida")
    ws = week_start_of(today)
    if target > today:
        raise HTTPException(status_code=400, detail="No puedes reportar un día que aún no ha llegado")
    if not (ws <= target <= ws + timedelta(days=6)):
        raise HTTPException(status_code=400, detail="Solo puedes reportar días de la semana en curso (miércoles a martes)")
    existing = await db.entries.find_one({"worker_id": user["id"], "date": target.isoformat()})
    if existing:
        raise HTTPException(status_code=409, detail="Ya registraste tus minutos de ese día")
    max_bytes = 10 * 1024 * 1024
    data = await screenshot.read(max_bytes + 1)
    if len(data) > max_bytes:
        raise HTTPException(status_code=400, detail="La imagen es demasiado grande")
    sniffed = sniff_image(data)
    if not sniffed:
        raise HTTPException(status_code=400, detail="El archivo no es una imagen válida (JPG, PNG, WEBP, GIF o HEIC)")
    content_type, ext = sniffed
    path = f"{APP_NAME}/uploads/{user['id']}/{uuid.uuid4()}.{ext}"
    try:
        result = await asyncio.to_thread(put_object, path, data, content_type)
    except Exception as e:
        logger.error(f"Storage upload failed: {e}")
        raise HTTPException(status_code=502, detail="No se pudo subir la captura")
    doc = {
        "id": str(uuid.uuid4()),
        "worker_id": user["id"],
        "date": target.isoformat(),
        "minutes": minutes,
        "screenshot_path": result["path"],
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.entries.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.delete("/entries/{entry_id}")
async def delete_my_entry(entry_id: str, user=Depends(get_current_user)):
    if user.get("role") != "worker":
        raise HTTPException(status_code=403, detail="Solo miembros")
    entry = await db.entries.find_one({"id": entry_id, "worker_id": user["id"]}, {"_id": 0})
    if not entry:
        raise HTTPException(status_code=404, detail="Registro no encontrado")
    ws = week_start_of(today_local())
    if not (ws.isoformat() <= entry["date"] <= (ws + timedelta(days=6)).isoformat()):
        raise HTTPException(status_code=400, detail="Solo puedes borrar reportes de la semana en curso")
    await db.entries.delete_one({"id": entry_id})
    return {"deleted": True}


@api_router.post("/me/weekly-report-done")
async def weekly_report_done(user=Depends(get_current_user)):
    if user.get("role") != "worker":
        raise HTTPException(status_code=403, detail="Solo miembros")
    prev_ws = (week_start_of(today_local()) - timedelta(days=7)).isoformat()
    await db.users.update_one({"id": user["id"]}, {"$addToSet": {"weekly_report_done": prev_ws}})
    return {"done": prev_ws}


def build_member_weeks(entries, paid_weeks, ws, today):
    week_starts = {week_start_of(date.fromisoformat(e["date"])) for e in entries} | {ws, ws + timedelta(days=7)}
    weeks = []
    for w0 in sorted(week_starts):
        w1 = w0 + timedelta(days=6)
        m = sum(e["minutes"] for e in entries if w0.isoformat() <= e["date"] <= w1.isoformat())
        if w0 == ws:
            status = "current"
        elif w0 > ws:
            status = "future"
        else:
            status = "past"
        weeks.append({
            "start": w0.isoformat(),
            "end": w1.isoformat(),
            "kgen_payday": (w0 + timedelta(days=12)).isoformat(),
            "bonus_payday": (w0 + timedelta(days=13)).isoformat(),
            "minutes": m,
            "qualifies": m > GOAL_MINUTES,
            "paid": w0.isoformat() in paid_weeks,
            "closed": w1 < today,
            "status": status,
        })
    return weeks


async def compute_streaks(user, entry_dates, today):
    streak = 0
    cursor = today if today.isoformat() in entry_dates else today - timedelta(days=1)
    while cursor.isoformat() in entry_dates:
        streak += 1
        cursor -= timedelta(days=1)
    best = 0
    run = 0
    prev = None
    for d in sorted(date.fromisoformat(x) for x in entry_dates):
        run = run + 1 if prev is not None and (d - prev).days == 1 else 1
        best = max(best, run)
        prev = d
    stored = int(user.get("best_streak") or 0)
    best = max(best, stored, streak)
    if best != stored:
        await db.users.update_one({"id": user["id"]}, {"$set": {"best_streak": best}})
    return streak, best


async def build_earnings(user_id, weeks, payments):
    pay_map = {p["week_start"]: p for p in payments}
    earnings = []
    for wk in weeks:
        if wk["status"] == "future":
            continue
        w0 = date.fromisoformat(wk["start"])
        rates = await rates_for_week(w0)
        p = pay_map.get(wk["start"])
        if p:
            bonus_total = p["amount"]
        else:
            wrows, _ = await build_week_table(w0)
            mine = next((x for x in wrows if x["id"] == user_id), None)
            bonus_total = mine["total"] if mine else 0
        earnings.append({
            "week_start": wk["start"],
            "label": fmt_short(w0),
            "kgen": round(wk["minutes"] / 60 * rates["kgen_rate"], 2),
            "bonus": round(bonus_total, 2),
            "minutes": wk["minutes"],
            "paid": bool(p),
            "current": wk["status"] == "current",
        })
    return earnings


def weekly_report_info(user, entries, ws, now):
    prev_ws = ws - timedelta(days=7)
    prev_we = prev_ws + timedelta(days=6)
    prev_minutes = sum(e["minutes"] for e in entries if prev_ws.isoformat() <= e["date"] <= prev_we.isoformat())
    notice_from = datetime.combine(ws, datetime.min.time(), tzinfo=TZ).replace(hour=12)
    return {
        "due": prev_minutes > 300 and now >= notice_from,
        "done": prev_ws.isoformat() in (user.get("weekly_report_done") or []),
        "prev_week_start": prev_ws.isoformat(),
        "prev_week_end": prev_we.isoformat(),
        "prev_minutes": prev_minutes,
    }


def week_summary(ws, we, week_minutes, me_row):
    row = me_row or {}
    return {
        "start": ws.isoformat(),
        "end": we.isoformat(),
        "payday": (ws + timedelta(days=13)).isoformat(),
        "kgen_payday": (ws + timedelta(days=12)).isoformat(),
        "minutes": week_minutes,
        "hours": round(week_minutes / 60, 2),
        "goal_minutes": GOAL_MINUTES,
        "remaining_minutes": max(0, GOAL_MINUTES - week_minutes),
        "qualifies": row.get("qualifies", False),
        "base": row.get("base", 0),
        "bonus_pct": row.get("bonus_pct", 0),
        "bonus": row.get("bonus", 0),
        "estimated_total": row.get("total", 0),
        "kgen": row.get("kgen", 0),
        "rank": row.get("rank"),
    }


@api_router.get("/me/dashboard")
async def me_dashboard(user=Depends(get_current_user)):
    if user.get("role") != "worker":
        raise HTTPException(status_code=403, detail="Solo miembros")
    today = today_local()
    now = now_local()
    ws = week_start_of(today)
    we = ws + timedelta(days=6)
    entries = await db.entries.find({"worker_id": user["id"]}, {"_id": 0}).sort("date", -1).to_list(1000)
    payments = await db.payments.find({"worker_id": user["id"]}, {"_id": 0}).to_list(1000)
    weeks = build_member_weeks(entries, {p["week_start"] for p in payments}, ws, today)
    week_minutes = sum(e["minutes"] for e in entries if ws.isoformat() <= e["date"] <= we.isoformat())
    today_entry = next((e for e in entries if e["date"] == today.isoformat()), None)
    rows, _ = await build_week_table(ws)
    me_row = next((r for r in rows if r["id"] == user["id"]), None)
    global_ranking = await build_global_ranking()
    my_global = next((g for g in global_ranking if g["id"] == user["id"]), None)
    streak, best_streak = await compute_streaks(user, {e["date"] for e in entries}, today)
    rates = await rates_for_week(ws)
    hours_now = week_minutes / 60
    return {
        "name": user["name"],
        "now": now.isoformat(),
        "today_date": today.isoformat(),
        "rates": rates,
        "estimate": {
            "kgen": round(hours_now * rates["kgen_rate"], 2),
            "bonus_if_goal": round(hours_now * rates["bonus_rate"], 2),
            "bonus_now": me_row["total"] if me_row else 0,
        },
        "earnings": await build_earnings(user["id"], weeks, payments),
        "weekly_report": weekly_report_info(user, entries, ws, now),
        "binance_pay_id": user.get("binance_pay_id") or "",
        "usdt_bep20_address": user.get("usdt_bep20_address") or "",
        "streak": streak,
        "on_fire": bool(entries) and entries[0]["minutes"] > 300,
        "best_streak": best_streak,
        "days": [
            {"date": (ws + timedelta(days=i)).isoformat(), "label": DAY_LABELS[i], "total": sum(e["minutes"] for e in entries if e["date"] == (ws + timedelta(days=i)).isoformat())}
            for i in range(7)
        ],
        "weeks": weeks,
        "day_minutes": {e["date"]: e["minutes"] for e in entries},
        "day_reviewed": [e["date"] for e in entries if e.get("reviewed")],
        "today": {
            "registered": today_entry is not None,
            "minutes": today_entry["minutes"] if today_entry else 0,
            "screenshot_path": today_entry["screenshot_path"] if today_entry else None,
        },
        "week": week_summary(ws, we, week_minutes, me_row),
        "global_rank": my_global["rank"] if my_global else None,
        "global_minutes": my_global["minutes"] if my_global else 0,
        "historical_minutes": int(user.get("historical_minutes") or 0),
        "weekly": [{"rank": r["rank"], "id": r["id"], "name": r["name"], "minutes": r["minutes"], "hours": r["hours"]} for r in rows],
        "global": global_ranking,
        "entries": entries[:30],
    }


@api_router.get("/files/{path:path}")
async def serve_file(path: str, token: str = Query(...)):
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
    except jwt.PyJWTError:
        raise HTTPException(status_code=401, detail="Token inválido")
    if payload.get("scope") != "file":
        raise HTTPException(status_code=401, detail="Token inválido")
    parts = path.split("/")
    if ".." in parts or "" in parts or parts[0] != APP_NAME:
        raise HTTPException(status_code=404, detail="Archivo no encontrado")
    user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0})
    if not user or not user.get("active", True):
        raise HTTPException(status_code=401, detail="Usuario no encontrado")
    is_owner = len(parts) >= 4 and parts[1] == "uploads" and parts[2] == user["id"]
    if user.get("role") != "admin" and not is_owner:
        raise HTTPException(status_code=404, detail="Archivo no encontrado")
    try:
        data, _ = await asyncio.to_thread(get_object, path)
    except Exception:
        raise HTTPException(status_code=404, detail="Archivo no encontrado")
    sniffed = sniff_image(data)
    if not sniffed:
        raise HTTPException(status_code=404, detail="Archivo no encontrado")
    return Response(
        content=data,
        media_type=sniffed[0],
        headers={
            "X-Content-Type-Options": "nosniff",
            "Content-Security-Policy": "default-src 'none'; sandbox",
            "Cache-Control": "private, max-age=300",
            "Content-Disposition": "inline",
        },
    )


@asynccontextmanager
async def lifespan(app_: FastAPI):
    await startup()
    yield
    client.close()


app = FastAPI(lifespan=lifespan)
app.include_router(api_router)

_cors = [o.strip() for o in (os.environ.get("CORS_ORIGINS") or os.environ.get("FRONTEND_URL") or "").split(",") if o.strip()]
if not _cors or "*" in _cors:
    logger.warning("CORS abierto (*). Define CORS_ORIGINS con los dominios reales del frontend.")
    _cors = ["*"]
app.add_middleware(
    CORSMiddleware,
    allow_credentials=False,  # se usa Bearer token, no cookies
    allow_origins=_cors,
    allow_methods=["GET", "POST", "PUT", "PATCH", "DELETE", "OPTIONS"],
    allow_headers=["Authorization", "Content-Type"],
)


async def startup():
    # Migración: PIN en texto plano -> hash + cifrado (primero quitar el índice único viejo)
    try:
        await db.users.drop_index("code_1")
    except Exception:
        pass
    async for u in db.users.find({"code": {"$exists": True}}, {"_id": 0, "id": 1, "code": 1}):
        c = str(u["code"])
        await db.users.update_one(
            {"id": u["id"]},
            {"$set": {"code_hash": hash_code(c), "code_enc": encrypt_code(c)}, "$unset": {"code": ""}},
        )
    await db.users.create_index("code_hash", unique=True, partialFilterExpression={"code_hash": {"$exists": True}})
    await db.entries.create_index([("worker_id", 1), ("date", 1)], unique=True)
    await db.payments.create_index([("worker_id", 1), ("week_start", 1)], unique=True)
    existing = await db.users.find_one({"role": "admin"})
    if not existing:
        if not (ADMIN_CODE and ADMIN_CODE.isdigit() and len(ADMIN_CODE) == 4):
            raise RuntimeError("No hay admin: define ADMIN_CODE (4 dígitos) para crearlo la primera vez")
        await db.users.insert_one({
            "id": str(uuid.uuid4()),
            "name": "Wuilber",
            "email": ADMIN_EMAIL,
            "code_hash": hash_code(ADMIN_CODE),
            "code_enc": encrypt_code(ADMIN_CODE),
            "role": "admin",
            "active": True,
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
        logger.info("Admin Wuilber creado")
    else:
        await db.users.update_one(
            {"role": "admin"},
            {"$set": {"name": "Wuilber", "active": True, **({"email": ADMIN_EMAIL} if ADMIN_EMAIL else {})}},
        )
    try:
        await asyncio.to_thread(init_storage)
        logger.info("Storage inicializado")
    except Exception as e:
        logger.error(f"Storage init falló: {e}")
